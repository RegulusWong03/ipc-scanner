"""Phase 4 导入导出增强测试"""

import sys
import os
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from device.models import Device, DeviceStatus, DeviceType
from export.handlers import (
    export_excel, export_csv, import_excel, import_csv,
    generate_template, validate_import_data,
    _is_valid_mac, _is_valid_ip,
)


class TestValidation:
    """导入数据校验测试"""

    def test_is_valid_mac_colon(self):
        assert _is_valid_mac("aa:bb:cc:dd:ee:ff") is True
        assert _is_valid_mac("AA:BB:CC:DD:EE:FF") is True
        assert _is_valid_mac("00:11:22:33:44:55") is True

    def test_is_valid_mac_dash(self):
        assert _is_valid_mac("aa-bb-cc-dd-ee-ff") is True

    def test_is_valid_mac_no_separator(self):
        assert _is_valid_mac("aabbccddeeff") is True

    def test_is_valid_mac_invalid(self):
        assert _is_valid_mac("") is False
        assert _is_valid_mac("aa:bb:cc") is False
        assert _is_valid_mac("gg:hh:ii:jj:kk:ll") is False
        assert _is_valid_mac("zzzz") is False
        assert _is_valid_mac("aa:bb:cc:dd:ee") is False

    def test_is_valid_ip_valid(self):
        assert _is_valid_ip("192.168.1.1") is True
        assert _is_valid_ip("10.0.0.1") is True
        assert _is_valid_ip("255.255.255.255") is True
        assert _is_valid_ip("0.0.0.0") is True

    def test_is_valid_ip_invalid(self):
        assert _is_valid_ip("") is False
        assert _is_valid_ip("256.0.0.1") is False
        assert _is_valid_ip("192.168.1") is False
        assert _is_valid_ip("192.168.1.1.1") is False
        assert _is_valid_ip("abc.def.ghi.jkl") is False

    def test_validate_import_data_valid(self):
        """测试有效数据通过校验"""
        rows = [
            {"ip": "192.168.1.100", "mac": "aa:bb:cc:dd:ee:ff", "brand": "hikvision"},
            {"ip": "192.168.1.101", "mac": "aa:bb:cc:dd:ee:01", "brand": "dahua"},
        ]
        valid, errors = validate_import_data(rows)
        assert len(valid) == 2
        assert len(errors) == 0

    def test_validate_import_data_missing_mac(self):
        """测试 MAC 为空报错"""
        rows = [{"ip": "192.168.1.100", "mac": "", "brand": "hikvision"}]
        valid, errors = validate_import_data(rows)
        assert len(valid) == 0
        assert len(errors) == 1
        assert "MAC 为空" in errors[0]["_errors"]

    def test_validate_import_data_invalid_mac(self):
        """测试 MAC 格式错误"""
        rows = [{"ip": "192.168.1.100", "mac": "invalid", "brand": "hikvision"}]
        valid, errors = validate_import_data(rows)
        assert len(valid) == 0
        assert len(errors) == 1
        assert "MAC 格式错误" in errors[0]["_errors"]

    def test_validate_import_data_missing_ip(self):
        """测试 IP 为空报错"""
        rows = [{"ip": "", "mac": "aa:bb:cc:dd:ee:ff", "brand": "hikvision"}]
        valid, errors = validate_import_data(rows)
        assert len(valid) == 0
        assert len(errors) == 1
        assert "IP 为空" in errors[0]["_errors"]

    def test_validate_import_data_invalid_ip(self):
        """测试 IP 格式错误"""
        rows = [{"ip": "999.999.999.999", "mac": "aa:bb:cc:dd:ee:ff"}]
        valid, errors = validate_import_data(rows)
        assert len(valid) == 0
        assert len(errors) == 1
        assert "IP 格式错误" in errors[0]["_errors"]

    def test_validate_import_data_invalid_subnet(self):
        """测试子网掩码格式错误"""
        rows = [{
            "ip": "192.168.1.100", "mac": "aa:bb:cc:dd:ee:ff",
            "subnet_mask": "invalid",
        }]
        valid, errors = validate_import_data(rows)
        assert len(valid) == 0
        assert len(errors) == 1
        assert "子网掩码格式错误" in errors[0]["_errors"]

    def test_validate_import_data_mixed(self):
        """测试混合数据（有效+无效）"""
        rows = [
            {"ip": "192.168.1.100", "mac": "aa:bb:cc:dd:ee:ff"},  # 有效
            {"ip": "192.168.1.101", "mac": ""},                     # 无效
            {"ip": "invalid", "mac": "aa:bb:cc:dd:ee:01"},         # 无效
            {"ip": "10.0.0.1", "mac": "11:22:33:44:55:66"},        # 有效
        ]
        valid, errors = validate_import_data(rows)
        assert len(valid) == 2
        assert len(errors) == 2


class TestExportExcel:
    """Excel 导出测试"""

    def test_export_excel_basic(self, tmp_path):
        """测试基本导出"""
        devices = [
            Device(mac="aa:bb:cc:dd:ee:ff", ip="192.168.1.100", brand="hikvision",
                   model="DS-2CD2142FWD", device_type=DeviceType.IPC, channels=1),
        ]
        filepath = str(tmp_path / "test.xlsx")
        export_excel(devices, filepath, styled=False)
        assert os.path.exists(filepath)

        # 验证能重新导入
        rows = import_excel(filepath)
        assert len(rows) == 1
        assert rows[0]["mac"] == "aa:bb:cc:dd:ee:ff"
        assert rows[0]["ip"] == "192.168.1.100"
        assert rows[0]["brand"] == "hikvision"

    def test_export_excel_styled(self, tmp_path):
        """测试带样式导出"""
        devices = [
            Device(mac="aa:bb:cc:dd:ee:ff", ip="192.168.1.100", brand="hikvision"),
            Device(mac="11:22:33:44:55:66", ip="192.168.1.101", brand="dahua"),
        ]
        filepath = str(tmp_path / "styled.xlsx")
        export_excel(devices, filepath, styled=True)
        assert os.path.exists(filepath)

        # 验证样式已应用
        from openpyxl import load_workbook
        wb = load_workbook(filepath)
        ws = wb.active
        # 检查表头样式
        header_cell = ws.cell(row=1, column=1)
        assert header_cell.font.bold is True
        assert header_cell.font.color.rgb == "00FFFFFF"
        # 检查冻结首行
        assert ws.freeze_panes == "A2"
        # 检查筛选器
        assert ws.auto_filter.ref is not None

    def test_export_excel_empty(self, tmp_path):
        """测试空设备列表导出"""
        filepath = str(tmp_path / "empty.xlsx")
        export_excel([], filepath)
        assert os.path.exists(filepath)

        rows = import_excel(filepath)
        assert len(rows) == 0

    def test_export_csv(self, tmp_path):
        """测试 CSV 导出"""
        devices = [
            Device(mac="aa:bb:cc:dd:ee:ff", ip="192.168.1.100", brand="hikvision"),
        ]
        filepath = str(tmp_path / "test.csv")
        export_csv(devices, filepath)
        assert os.path.exists(filepath)

        rows = import_csv(filepath)
        assert len(rows) == 1
        assert rows[0]["mac"] == "aa:bb:cc:dd:ee:ff"


class TestTemplate:
    """模板生成测试"""

    def test_generate_template(self, tmp_path):
        """测试模板生成"""
        filepath = str(tmp_path / "template.xlsx")
        generate_template(filepath)
        assert os.path.exists(filepath)

        from openpyxl import load_workbook
        wb = load_workbook(filepath)
        ws = wb.active
        assert ws.title == "设备清单模板"

        # 检查表头
        headers = [ws.cell(row=1, column=i).value for i in range(1, 15)]
        assert "IP" in headers
        assert "MAC" in headers

        # 检查示例行
        assert ws.cell(row=2, column=1).value == "192.168.1.100"
        assert ws.cell(row=2, column=2).value == "aa:bb:cc:dd:ee:ff"


class TestImportExportRoundtrip:
    """导入导出往返测试"""

    def test_excel_roundtrip(self, tmp_path):
        """测试 Excel 导出再导入数据一致"""
        devices = [
            Device(mac="aa:bb:cc:dd:ee:ff", ip="192.168.1.100",
                   brand="hikvision", model="DS-2CD2142FWD",
                   firmware_version="V5.4.3", serial_number="SN001",
                   channels=1, subnet_mask="255.255.255.0",
                   gateway="192.168.1.1", dhcp=False,
                   group="1F-大厅", note="前台摄像头"),
            Device(mac="11:22:33:44:55:66", ip="192.168.1.101",
                   brand="dahua", model="IPC-HDW4431C",
                   device_type=DeviceType.IPC, channels=1),
        ]

        filepath = str(tmp_path / "roundtrip.xlsx")
        export_excel(devices, filepath)
        rows = import_excel(filepath)

        assert len(rows) == 2
        assert rows[0]["mac"] == "aa:bb:cc:dd:ee:ff"
        assert rows[0]["ip"] == "192.168.1.100"
        assert rows[0]["brand"] == "hikvision"
        assert rows[0]["model"] == "DS-2CD2142FWD"
        assert rows[0]["group"] == "1F-大厅"
        assert rows[0]["note"] == "前台摄像头"

        assert rows[1]["mac"] == "11:22:33:44:55:66"
        assert rows[1]["brand"] == "dahua"

    def test_csv_roundtrip(self, tmp_path):
        """测试 CSV 导出再导入数据一致"""
        devices = [
            Device(mac="aa:bb:cc:dd:ee:ff", ip="192.168.1.100",
                   brand="hikvision", model="DS-2CD2142FWD"),
        ]

        filepath = str(tmp_path / "roundtrip.csv")
        export_csv(devices, filepath)
        rows = import_csv(filepath)

        assert len(rows) == 1
        assert rows[0]["mac"] == "aa:bb:cc:dd:ee:ff"
        assert rows[0]["ip"] == "192.168.1.100"


class TestImportDialogUI:
    """导入对话框 UI 测试"""

    def test_import_dialog_creation(self):
        """测试对话框创建"""
        os.environ["QT_QPA_PLATFORM"] = "offscreen"
        from PyQt6.QtWidgets import QApplication
        app = QApplication.instance() or QApplication(sys.argv)
        from ui.import_dialog import ImportDialog

        rows = [
            {"ip": "192.168.1.100", "mac": "aa:bb:cc:dd:ee:ff", "brand": "hikvision"},
        ]
        existing_macs = set()
        dialog = ImportDialog(rows, existing_macs)

        assert dialog is not None
        assert dialog.preview_table is not None
        assert dialog.radio_skip.isChecked()

    def test_import_dialog_with_existing(self):
        """测试对话框显示已存在设备"""
        os.environ["QT_QPA_PLATFORM"] = "offscreen"
        from PyQt6.QtWidgets import QApplication
        app = QApplication.instance() or QApplication(sys.argv)
        from ui.import_dialog import ImportDialog

        rows = [
            {"ip": "192.168.1.100", "mac": "aa:bb:cc:dd:ee:ff", "brand": "hikvision"},
            {"ip": "192.168.1.101", "mac": "11:22:33:44:55:66", "brand": "dahua"},
        ]
        existing_macs = {"aa:bb:cc:dd:ee:ff"}  # 第一个已存在
        dialog = ImportDialog(rows, existing_macs)

        # 预览表格应有 2 行
        assert dialog.preview_table.rowCount() == 2

    def test_import_dialog_get_result(self):
        """测试获取导入结果"""
        os.environ["QT_QPA_PLATFORM"] = "offscreen"
        from PyQt6.QtWidgets import QApplication
        app = QApplication.instance() or QApplication(sys.argv)
        from ui.import_dialog import ImportDialog

        rows = [
            {"ip": "192.168.1.100", "mac": "aa:bb:cc:dd:ee:ff", "brand": "hikvision"},
        ]
        dialog = ImportDialog(rows, set())

        # 默认策略是 skip
        valid, strategy = dialog.get_result()
        assert strategy == ImportDialog.SKIP_EXISTING
        assert len(valid) == 1


class TestMainWindowPhase4:
    """主窗口 Phase 4 集成测试"""

    def test_main_window_toolbar_actions(self):
        """测试工具栏包含导入/导出/模板按钮"""
        os.environ["QT_QPA_PLATFORM"] = "offscreen"
        from PyQt6.QtWidgets import QApplication
        app = QApplication.instance() or QApplication(sys.argv)
        from ui.main_window import MainWindow

        window = MainWindow()

        assert window.import_action is not None
        assert window.export_action is not None
        assert window.export_selected_action is not None
        assert window.template_action is not None

    def test_device_table_multiselect(self):
        """测试设备表格支持多选"""
        os.environ["QT_QPA_PLATFORM"] = "offscreen"
        from PyQt6.QtWidgets import QApplication, QAbstractItemView
        app = QApplication.instance() or QApplication(sys.argv)
        from ui.device_table import DeviceTable

        table = DeviceTable()
        assert table.table.selectionMode() == QAbstractItemView.SelectionMode.ExtendedSelection

    def test_device_table_get_methods(self):
        """测试设备表格获取选中/过滤设备方法"""
        os.environ["QT_QPA_PLATFORM"] = "offscreen"
        from PyQt6.QtWidgets import QApplication
        app = QApplication.instance() or QApplication(sys.argv)
        from ui.device_table import DeviceTable
        from device.models import Device

        table = DeviceTable()
        devices = [
            Device(mac="aa:bb:cc:dd:ee:ff", ip="192.168.1.100"),
            Device(mac="11:22:33:44:55:66", ip="192.168.1.101"),
        ]
        table.update_devices(devices)

        # 未选择时返回空
        assert table.get_selected_devices() == []
        # 过滤后返回全部
        assert len(table.get_filtered_devices()) == 2


if __name__ == "__main__":
    import pytest
    sys.exit(pytest.main([__file__, "-v"]))
